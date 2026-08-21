from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import logging
import uuid
import io
from datetime import datetime, timezone, timedelta, date
from typing import List, Optional, Literal
from pydantic import BaseModel, Field
import bcrypt
import jwt
from bson import ObjectId

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

JWT_ALGORITHM = "HS256"
PRICE_TYPES = ["grosir", "so", "retail"]

# ----------------------- Auth helpers -----------------------

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))

def get_jwt_secret() -> str:
    return os.environ["JWT_SECRET"]

def create_access_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email,
               "exp": datetime.now(timezone.utc) + timedelta(hours=12), "type": "access"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

def create_refresh_token(user_id: str) -> str:
    payload = {"sub": user_id, "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "refresh"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

def set_auth_cookies(response: Response, access: str, refresh: str):
    response.set_cookie(key="access_token", value=access, httponly=True, secure=True,
                        samesite="none", max_age=43200, path="/")
    response.set_cookie(key="refresh_token", value=refresh, httponly=True, secure=True,
                        samesite="none", max_age=604800, path="/")

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Token tidak valid")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="User tidak ditemukan")
        user["_id"] = str(user["_id"])
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token kedaluwarsa")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")

# ----------------------- Models -----------------------

class LoginInput(BaseModel):
    email: str
    password: str

class UpdateNameInput(BaseModel):
    name: str

class UpdateEmailInput(BaseModel):
    new_email: str
    current_password: str

class UpdatePasswordInput(BaseModel):
    current_password: str
    new_password: str

class Variant(BaseModel):
    label: str
    harga_grosir: float = 0
    harga_so: float = 0
    harga_retail: float = 0

class ProductInput(BaseModel):
    nama: str
    variants: List[Variant] = []

class CustomerInput(BaseModel):
    nama: str
    telepon: str = ""
    alamat: str = ""
    default_price_type: Literal["grosir", "so", "retail"] = "retail"

class TransactionItem(BaseModel):
    product_id: str
    product_nama: str
    variant_label: str
    price_type: str
    harga_satuan: float
    qty: float
    subtotal: float

class TransactionInput(BaseModel):
    tanggal: str  # YYYY-MM-DD
    customer_id: str
    customer_nama: str
    items: List[TransactionItem]
    status: Literal["lunas", "belum_lunas"] = "lunas"
    catatan: str = ""

# ----------------------- Auth routes -----------------------

MAX_ATTEMPTS = 5
LOCKOUT_MINUTES = 15

@api_router.post("/auth/login")
async def login(data: LoginInput, request: Request, response: Response):
    email = data.email.lower().strip()
    ip = request.client.host if request.client else "unknown"
    identifier = f"{ip}:{email}"

    now = datetime.now(timezone.utc)
    rec = await db.login_attempts.find_one({"identifier": identifier})
    if rec and rec.get("count", 0) >= MAX_ATTEMPTS:
        last = rec.get("last")
        if isinstance(last, str):
            last = datetime.fromisoformat(last)
        if last and now - last < timedelta(minutes=LOCKOUT_MINUTES):
            raise HTTPException(status_code=429,
                detail=f"Terlalu banyak percobaan. Coba lagi dalam {LOCKOUT_MINUTES} menit.")

    user = await db.users.find_one({"email": email})
    if not user or not verify_password(data.password, user["password_hash"]):
        await db.login_attempts.update_one(
            {"identifier": identifier},
            {"$inc": {"count": 1}, "$set": {"last": now.isoformat()}},
            upsert=True)
        raise HTTPException(status_code=401, detail="Email atau password salah")

    await db.login_attempts.delete_one({"identifier": identifier})
    uid = str(user["_id"])
    access = create_access_token(uid, email)
    refresh = create_refresh_token(uid)
    set_auth_cookies(response, access, refresh)
    return {"id": uid, "email": email, "name": user.get("name", "Admin"), "role": user.get("role", "admin")}

@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    response.delete_cookie("refresh_token", path="/")
    return {"message": "Logout berhasil"}

@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return {"id": user["_id"], "email": user["email"], "name": user.get("name", "Admin"), "role": user.get("role", "admin")}

@api_router.put("/auth/profile/name")
async def update_name(data: UpdateNameInput, user: dict = Depends(get_current_user)):
    name = data.name.strip()
    if not name:
        raise HTTPException(status_code=422, detail="Nama tidak boleh kosong")
    await db.users.update_one({"_id": ObjectId(user["_id"])}, {"$set": {"name": name}})
    return {"id": user["_id"], "email": user["email"], "name": name, "role": user.get("role", "admin")}

@api_router.put("/auth/profile/email")
async def update_email(data: UpdateEmailInput, user: dict = Depends(get_current_user)):
    full = await db.users.find_one({"_id": ObjectId(user["_id"])})
    if not verify_password(data.current_password, full["password_hash"]):
        raise HTTPException(status_code=400, detail="Password saat ini salah")
    new_email = data.new_email.lower().strip()
    if not new_email or "@" not in new_email:
        raise HTTPException(status_code=422, detail="Email tidak valid")
    dup = await db.users.find_one({"email": new_email, "_id": {"$ne": ObjectId(user["_id"])}})
    if dup:
        raise HTTPException(status_code=400, detail="Email sudah digunakan")
    await db.users.update_one({"_id": ObjectId(user["_id"])}, {"$set": {"email": new_email}})
    return {"id": user["_id"], "email": new_email, "name": full.get("name", "Admin"), "role": full.get("role", "admin")}

@api_router.put("/auth/profile/password")
async def update_password_ep(data: UpdatePasswordInput, user: dict = Depends(get_current_user)):
    full = await db.users.find_one({"_id": ObjectId(user["_id"])})
    if not verify_password(data.current_password, full["password_hash"]):
        raise HTTPException(status_code=400, detail="Password saat ini salah")
    if len(data.new_password) < 6:
        raise HTTPException(status_code=422, detail="Password baru minimal 6 karakter")
    await db.users.update_one({"_id": ObjectId(user["_id"])}, {"$set": {"password_hash": hash_password(data.new_password)}})
    return {"message": "Password berhasil diperbarui"}

@api_router.post("/auth/refresh")
async def refresh_token(request: Request, response: Response):
    token = request.cookies.get("refresh_token")
    if not token:
        raise HTTPException(status_code=401, detail="Tidak ada refresh token")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "refresh":
            raise HTTPException(status_code=401, detail="Token tidak valid")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="User tidak ditemukan")
        access = create_access_token(str(user["_id"]), user["email"])
        response.set_cookie(key="access_token", value=access, httponly=True, secure=True,
                            samesite="none", max_age=43200, path="/")
        return {"message": "ok"}
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")

# ----------------------- Products -----------------------

@api_router.get("/products")
async def list_products(user: dict = Depends(get_current_user)):
    items = await db.products.find({}, {"_id": 0}).sort("nama", 1).to_list(1000)
    return items

@api_router.post("/products")
async def create_product(data: ProductInput, user: dict = Depends(get_current_user)):
    doc = {"id": str(uuid.uuid4()), "nama": data.nama,
           "variants": [v.model_dump() for v in data.variants],
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.products.insert_one(doc.copy())
    doc.pop("_id", None)
    return doc

@api_router.put("/products/{product_id}")
async def update_product(product_id: str, data: ProductInput, user: dict = Depends(get_current_user)):
    res = await db.products.update_one({"id": product_id}, {"$set": {
        "nama": data.nama, "variants": [v.model_dump() for v in data.variants]}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return await db.products.find_one({"id": product_id}, {"_id": 0})

@api_router.delete("/products/{product_id}")
async def delete_product(product_id: str, user: dict = Depends(get_current_user)):
    res = await db.products.delete_one({"id": product_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return {"message": "Produk dihapus"}

# ----------------------- Customers -----------------------

@api_router.get("/customers")
async def list_customers(user: dict = Depends(get_current_user)):
    items = await db.customers.find({}, {"_id": 0}).sort("nama", 1).to_list(1000)
    return items

@api_router.post("/customers")
async def create_customer(data: CustomerInput, user: dict = Depends(get_current_user)):
    doc = {"id": str(uuid.uuid4()), **data.model_dump(),
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.customers.insert_one(doc.copy())
    doc.pop("_id", None)
    return doc

@api_router.put("/customers/{customer_id}")
async def update_customer(customer_id: str, data: CustomerInput, user: dict = Depends(get_current_user)):
    res = await db.customers.update_one({"id": customer_id}, {"$set": data.model_dump()})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Customer tidak ditemukan")
    return await db.customers.find_one({"id": customer_id}, {"_id": 0})

@api_router.delete("/customers/{customer_id}")
async def delete_customer(customer_id: str, user: dict = Depends(get_current_user)):
    res = await db.customers.delete_one({"id": customer_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Customer tidak ditemukan")
    return {"message": "Customer dihapus"}

# ----------------------- Transactions -----------------------

def build_txn_query(date_from, date_to, customer_id, status):
    q = {}
    if date_from and date_to:
        q["tanggal"] = {"$gte": date_from, "$lte": date_to}
    elif date_from:
        q["tanggal"] = {"$gte": date_from}
    elif date_to:
        q["tanggal"] = {"$lte": date_to}
    if customer_id:
        q["customer_id"] = customer_id
    if status:
        q["status"] = status
    return q

@api_router.get("/transactions")
async def list_transactions(date_from: Optional[str] = None, date_to: Optional[str] = None,
                            customer_id: Optional[str] = None, status: Optional[str] = None,
                            user: dict = Depends(get_current_user)):
    q = build_txn_query(date_from, date_to, customer_id, status)
    items = await db.transactions.find(q, {"_id": 0}).sort("tanggal", -1).to_list(5000)
    return items

@api_router.post("/transactions")
async def create_transaction(data: TransactionInput, user: dict = Depends(get_current_user)):
    total = sum(i.subtotal for i in data.items)
    doc = {"id": str(uuid.uuid4()), "tanggal": data.tanggal, "customer_id": data.customer_id,
           "customer_nama": data.customer_nama, "items": [i.model_dump() for i in data.items],
           "total": total, "status": data.status, "catatan": data.catatan,
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.transactions.insert_one(doc.copy())
    doc.pop("_id", None)
    return doc

@api_router.put("/transactions/{txn_id}")
async def update_transaction(txn_id: str, data: TransactionInput, user: dict = Depends(get_current_user)):
    total = sum(i.subtotal for i in data.items)
    res = await db.transactions.update_one({"id": txn_id}, {"$set": {
        "tanggal": data.tanggal, "customer_id": data.customer_id, "customer_nama": data.customer_nama,
        "items": [i.model_dump() for i in data.items], "total": total,
        "status": data.status, "catatan": data.catatan}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")
    return await db.transactions.find_one({"id": txn_id}, {"_id": 0})

@api_router.patch("/transactions/{txn_id}/status")
async def update_status(txn_id: str, body: dict, user: dict = Depends(get_current_user)):
    status = body.get("status")
    if status not in ("lunas", "belum_lunas"):
        raise HTTPException(status_code=400, detail="Status tidak valid")
    await db.transactions.update_one({"id": txn_id}, {"$set": {"status": status}})
    return {"message": "Status diperbarui"}

@api_router.delete("/transactions/{txn_id}")
async def delete_transaction(txn_id: str, user: dict = Depends(get_current_user)):
    res = await db.transactions.delete_one({"id": txn_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")
    return {"message": "Transaksi dihapus"}

# ----------------------- Dashboard -----------------------

@api_router.get("/dashboard/stats")
async def dashboard_stats(user: dict = Depends(get_current_user)):
    txns = await db.transactions.find({}, {"_id": 0}).to_list(10000)
    total_omset = sum(t["total"] for t in txns)
    total_txn = len(txns)
    sisa_tagihan = sum(t["total"] for t in txns if t["status"] == "belum_lunas")
    total_customer = await db.customers.count_documents({})

    # omset trend by date (last 30 days that exist)
    by_date = {}
    for t in txns:
        by_date[t["tanggal"]] = by_date.get(t["tanggal"], 0) + t["total"]
    trend = [{"tanggal": k, "omset": v} for k, v in sorted(by_date.items())][-30:]

    # top products by omset
    prod = {}
    prod_qty = {}
    for t in txns:
        for it in t["items"]:
            key = f'{it["product_nama"]} {it["variant_label"]}'
            prod[key] = prod.get(key, 0) + it["subtotal"]
            prod_qty[key] = prod_qty.get(key, 0) + it["qty"]
    top_products = sorted(
        [{"nama": k, "omset": v, "qty": prod_qty[k]} for k, v in prod.items()],
        key=lambda x: x["omset"], reverse=True)[:8]

    # outstanding per customer
    outstanding = {}
    for t in txns:
        if t["status"] == "belum_lunas":
            outstanding[t["customer_nama"]] = outstanding.get(t["customer_nama"], 0) + t["total"]
    outstanding_list = sorted(
        [{"customer": k, "sisa": v} for k, v in outstanding.items()],
        key=lambda x: x["sisa"], reverse=True)

    return {"total_omset": total_omset, "total_txn": total_txn, "sisa_tagihan": sisa_tagihan,
            "total_customer": total_customer, "trend": trend, "top_products": top_products,
            "outstanding": outstanding_list}

@api_router.get("/rekap/monthly")
async def rekap_monthly(year: int, month: int, user: dict = Depends(get_current_user)):
    if month < 1 or month > 12:
        raise HTTPException(status_code=422, detail="Bulan harus antara 1 dan 12")
    if year < 2000 or year > 2100:
        raise HTTPException(status_code=422, detail="Tahun tidak valid")
    prefix = f"{year:04d}-{month:02d}"
    txns = await db.transactions.find({"tanggal": {"$regex": f"^{prefix}"}}, {"_id": 0}).to_list(10000)
    total_omset = sum(t["total"] for t in txns)
    total_txn = len(txns)
    sisa = sum(t["total"] for t in txns if t["status"] == "belum_lunas")
    lunas = total_omset - sisa

    prod = {}
    for t in txns:
        for it in t["items"]:
            key = f'{it["product_nama"]} {it["variant_label"]}'
            if key not in prod:
                prod[key] = {"nama": key, "qty": 0, "omset": 0}
            prod[key]["qty"] += it["qty"]
            prod[key]["omset"] += it["subtotal"]
    per_product = sorted(prod.values(), key=lambda x: x["omset"], reverse=True)

    cust = {}
    for t in txns:
        if t["customer_nama"] not in cust:
            cust[t["customer_nama"]] = {"customer": t["customer_nama"], "omset": 0, "sisa": 0, "txn": 0}
        cust[t["customer_nama"]]["omset"] += t["total"]
        cust[t["customer_nama"]]["txn"] += 1
        if t["status"] == "belum_lunas":
            cust[t["customer_nama"]]["sisa"] += t["total"]
    per_customer = sorted(cust.values(), key=lambda x: x["omset"], reverse=True)

    daily = {}
    for t in txns:
        daily[t["tanggal"]] = daily.get(t["tanggal"], 0) + t["total"]
    per_day = [{"tanggal": k, "omset": v} for k, v in sorted(daily.items())]

    return {"total_omset": total_omset, "total_txn": total_txn, "sisa_tagihan": sisa, "lunas": lunas,
            "per_product": per_product, "per_customer": per_customer, "per_day": per_day}

# ----------------------- Export -----------------------

def rupiah(n):
    return "Rp " + f"{int(round(n)):,}".replace(",", ".")

async def fetch_txns(date_from, date_to, customer_id, status):
    q = build_txn_query(date_from, date_to, customer_id, status)
    return await db.transactions.find(q, {"_id": 0}).sort("tanggal", 1).to_list(10000)

@api_router.get("/export/excel")
async def export_excel(date_from: Optional[str] = None, date_to: Optional[str] = None,
                       customer_id: Optional[str] = None, status: Optional[str] = None,
                       user: dict = Depends(get_current_user)):
    txns = await fetch_txns(date_from, date_to, customer_id, status)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Penjualan"
    header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    headers = ["Tanggal", "Customer", "Produk", "Varian", "Tipe Harga", "Harga Satuan", "Qty", "Subtotal", "Status"]
    ws.append(headers)
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
    for t in txns:
        for it in t["items"]:
            ws.append([t["tanggal"], t["customer_nama"], it["product_nama"], it["variant_label"],
                       it["price_type"], it["harga_satuan"], it["qty"], it["subtotal"],
                       "Lunas" if t["status"] == "lunas" else "Belum Lunas"])
    total = sum(t["total"] for t in txns)
    ws.append([])
    ws.append(["", "", "", "", "", "", "", "TOTAL OMSET", total])
    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=rekap_penjualan.xlsx"})

@api_router.get("/export/pdf")
async def export_pdf(date_from: Optional[str] = None, date_to: Optional[str] = None,
                     customer_id: Optional[str] = None, status: Optional[str] = None,
                     user: dict = Depends(get_current_user)):
    txns = await fetch_txns(date_from, date_to, customer_id, status)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=18, bottomMargin=18,
                            leftMargin=18, rightMargin=18)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#0f172a"))
    elems = [Paragraph("Rekap Laporan Penjualan", title_style)]
    subtitle = f"Periode: {date_from or 'Semua'} s/d {date_to or 'Semua'}"
    elems.append(Paragraph(subtitle, styles["Normal"]))
    elems.append(Spacer(1, 10))
    data = [["Tanggal", "Customer", "Produk", "Varian", "Tipe", "Harga", "Qty", "Subtotal", "Status"]]
    for t in txns:
        for it in t["items"]:
            data.append([t["tanggal"], t["customer_nama"], it["product_nama"], it["variant_label"],
                         it["price_type"], rupiah(it["harga_satuan"]), str(int(it["qty"])),
                         rupiah(it["subtotal"]), "Lunas" if t["status"] == "lunas" else "Belum Lunas"])
    total = sum(t["total"] for t in txns)
    data.append(["", "", "", "", "", "", "", "TOTAL", rupiah(total)])
    tbl = Table(data, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f1f5f9")]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    elems.append(tbl)
    doc.build(elems)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=rekap_penjualan.pdf"})

# ----------------------- Seed -----------------------

SEED_PRODUCTS = [
    ("CHEERS ALKALINE", ["1200", "550", "330", "230", "GALON"]),
    ("CHEERS REGULAR", ["1500", "600", "330", "220", "GALON"]),
    ("VEMA", ["1500", "600", "220", "GALON"]),
]
DEFAULT_PRICES = {
    "1200": (45000, 47000, 50000), "1500": (40000, 42000, 45000),
    "600": (30000, 32000, 35000), "550": (30000, 32000, 35000),
    "330": (25000, 27000, 30000), "230": (20000, 22000, 25000),
    "220": (20000, 22000, 24000), "GALON": (18000, 19000, 20000),
}

@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("identifier", unique=True)
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@example.com").lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "admin123")
    # Seed admin only when no user exists yet. Do NOT overwrite an existing
    # user's email/password on restart so self-service profile edits persist.
    if await db.users.count_documents({}) == 0:
        await db.users.insert_one({"email": admin_email, "password_hash": hash_password(admin_password),
                                   "name": "Admin", "role": "admin",
                                   "created_at": datetime.now(timezone.utc).isoformat()})
        logger.info("Admin dibuat")

    if await db.products.count_documents({}) == 0:
        for nama, labels in SEED_PRODUCTS:
            variants = []
            for lb in labels:
                g, s, r = DEFAULT_PRICES.get(lb, (0, 0, 0))
                variants.append({"label": lb, "harga_grosir": g, "harga_so": s, "harga_retail": r})
            await db.products.insert_one({"id": str(uuid.uuid4()), "nama": nama, "variants": variants,
                                          "created_at": datetime.now(timezone.utc).isoformat()})
        logger.info("Produk awal dibuat")

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[os.environ.get("FRONTEND_URL", "http://localhost:3000"), "http://localhost:3000"],
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
